import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT, FILL_TEST = sys.argv[1], len(sys.argv) > 2
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10); BLK = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
GREEN = Font(name=F, color="008000", size=10); H1 = Font(name=F, size=16, bold=True); H2 = Font(name=F, size=11, bold=True, color="FFFFFF")
NOTE = Font(name=F, size=9, italic=True, color="666666")
YEL = PatternFill("solid", fgColor="FFFF00"); DARK = PatternFill("solid", fgColor="14100D"); GREY = PatternFill("solid", fgColor="EFEBE3"); RUST = PatternFill("solid", fgColor="F6D9D1")
THIN = Side(style="thin", color="BBBBBB"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AUD = '"A$"#,##0.00;("A$"#,##0.00);"-"'; NUM = '#,##0.00;(#,##0.00);"-"'; PCT = '0.0%;(0.0%);"-"'

wb = Workbook(); wb.properties.creator = "Sanchez Custom Boxing"; wb.properties.lastModifiedBy = "Sanchez Custom Boxing"; wb.properties.title = "Price breakdown"

def put(ws, ref, v, font=BLK, fmt=None, fill=None, bold=False, align=None, border=True):
    c = ws[ref]; c.value = v; c.font = BOLD if bold else font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if border: c.border = BOX
    return c
def inp(ws, ref, v=None, fmt=None):  # a cell the reader fills in
    return put(ws, ref, v, font=BLUE, fmt=fmt, fill=YEL)
def band(ws, row, text, cols="ABCDEF"):
    for col in cols: put(ws, f"{col}{row}", None, fill=DARK)
    ws[f"A{row}"].value = text; ws[f"A{row}"].font = H2
def widths(ws, w):
    for col, n in w.items(): ws.column_dimensions[col].width = n

# ---------------- Read me
rm = wb.active; rm.title = "Read me"; widths(rm, {"A": 3, "B": 110})
rm["B2"].value = "Sanchez Custom Boxing: price breakdown"; rm["B2"].font = H1
lines = [
 ("What this is", True),
 ("One sheet per product. Each takes one real, finished piece (the example named at the top) and breaks its price down to the last cost: materials, labour, factory overhead, rejects, packaging, shipping, tax, payment fees and Michael's share. What is left is what Sanchez keeps.", False),
 ("", False),
 ("How to fill it in", True),
 ("1. Settings first: exchange rate, labour rate, monthly overhead, tax, fees and shares. They feed every product sheet.", False),
 ("2. Then each product sheet, top to bottom. Yellow cells with blue text are yours to fill in. Everything else is a formula: do not type over it.", False),
 ("3. Every cost line has its own currency (THB or AUD). The sheet converts THB to AUD with the rate in Settings.", False),
 ("4. Type the selling price last. The sheet shows the profit and margin for both shipping routes, the break-even price, and the price that hits the target margin.", False),
 ("5. Summary puts all five products side by side.", False),
 ("", False),
 ("The two shipping routes", True),
 ("Route A, via our address: the factory ships to us, we send it on to the customer. This is how it works today.", False),
 ("Route B, direct: the factory ships straight to the customer.", False),
 ("", False),
 ("Colours", True),
 ("Yellow cell, blue text: a number you type. Black: a formula. Green: a number carried from another sheet.", False),
 ("", False),
 ("Format example (not a real figure, do not copy it)", True),
 ("A materials line reads: Item 'Leather, outer' | Qty 0.35 | Unit 'sq m' | Unit cost 1,200.00 | Currency THB. A labour line reads: Task 'Stitching' | Hours 1.50. Percentages are typed as 10%, not 0.1.", False),
 ("", False),
 ("Every number in this workbook must be Jesse's real figure. Nothing has been filled in for him: the item names are prompts only, and a blank line can be renamed or left empty.", True),
]
for i, (t, b) in enumerate(lines):
    c = rm[f"B{4+i}"]; c.value = t; c.font = BOLD if b else BLK; c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Settings
st = wb.create_sheet("Settings"); widths(st, {"A": 46, "B": 16, "C": 70})
st["A1"].value = "Settings: used by every product sheet"; st["A1"].font = H1
band(st, 3, "Setting", "ABC"); st["B3"].value = "Value"; st["B3"].font = H2; st["C3"].value = "What it is"; st["C3"].font = H2
S = [  # label, name, fmt, note
 ("Exchange rate: A$ for 1 Thai baht", "fx", '0.0000', "Source to be noted here by whoever fills it in (bank or card rate actually paid, with the date)."),
 ("Labour rate (THB per hour)", "rate", NUM, "What an hour of making costs, wages plus on-costs. Each labour line can override it."),
 ("Factory overhead per month (THB)", "oh", NUM, "Rent, power, machine upkeep, anything not tied to one piece."),
 ("Pieces made per month, all products", "pcs", '#,##0;-#,##0;"-"', "Overhead is spread evenly over this many pieces."),
 ("Rejects and rework allowance", "rej", PCT, "Share of materials and labour lost to pieces that fail QC. Type as a percentage."),
 ("GST or VAT charged to the customer", "gst", PCT, "Australia: 10% once registered. Type 0% if prices carry no tax."),
 ("Payment fee, percentage of the sale", "fee", PCT, "The card processor's percentage, from its own price list."),
 ("Payment fee, fixed amount per sale (A$)", "fix", AUD, "The card processor's fixed charge per transaction."),
 ("Michael's share of each sale", "share", PCT, "As agreed in the stage agreement. Taken on the price before GST."),
 ("Target margin for Sanchez", "target", PCT, "Profit as a share of the price before GST. Used for the 'price to hit target' line."),
]
REF = {}
for i, (lab, key, fmt, note) in enumerate(S):
    r = 4 + i; put(st, f"A{r}", lab); inp(st, f"B{r}", None, fmt); put(st, f"C{r}", note, font=NOTE); REF[key] = f"Settings!$B${r}"
r = 4 + len(S) + 1
put(st, f"A{r}", "Overhead per piece (A$)", bold=True); put(st, f"B{r}", f"=IF({REF['pcs'][9:]}>0,{REF['oh'][9:]}*{REF['fx'][9:]}/{REF['pcs'][9:]},0)", fmt=AUD); put(st, f"C{r}", "Monthly overhead, converted, divided by pieces per month.", font=NOTE)
REF["ohpp"] = f"Settings!$B${r}"
st.freeze_panes = "A4"

# ---------------- product sheets
PRODUCTS = [
 ("Focus mitts", "Brown focus mitts with stitched initials (the brown pair on the site)", "pair",
  ["Leather, outer", "Leather or vinyl, striking face", "Foam and padding", "Lining", "Thread", "Lacing", "Strap and fastening", "Sanchez badge", "Lettering or initials", "Adhesive", "", ""]),
 ("Gloves", "Example piece: to be chosen with Jesse (no photograph supplied yet)", "pair",
  ["Leather, outer", "Foam and padding", "Lining", "Thread", "Piping", "Laces or strap", "Sanchez badge", "Printing or embroidery", "Adhesive", "", "", ""]),
 ("Heavy bag", "Heavy bag in red, white and green (the tricolour bag on the site)", "bag",
  ["Persian vinyl", "HH-66 vinyl cement", "Thread", "Chain and swivel", "D-rings and top hardware", "Zip", "Filling (if sent filled)", "Screen print", "Sanchez label", "", "", ""]),
 ("Head guard", "Example piece: to be chosen with Jesse (no photograph supplied yet)", "piece",
  ["Leather, outer", "Foam and padding", "Lining", "Thread", "Strap, buckle or lacing", "Sanchez badge", "Printing or embroidery", "Adhesive", "", "", "", ""]),
 ("Groin guard", "Red and blue groin guard with white lettering (the pair on the site)", "piece",
  ["Leather, outer", "Foam and padding", "Cup or hard insert", "Lining", "Thread", "Strap and fastening", "Lettering", "Sanchez badge", "Adhesive", "", "", ""]),
]
TASKS = ["Cutting", "Stitching", "Assembly and gluing", "Printing or lettering", "Finishing and quality check", "Packing"]
dv_src = '"THB,AUD"'
cells = {}
for name, example, unit, mats in PRODUCTS:
    ws = wb.create_sheet(name); widths(ws, {"A": 44, "B": 11, "C": 12, "D": 14, "E": 11, "F": 16, "G": 3, "H": 52})
    dv = DataValidation(type="list", formula1=dv_src, allow_blank=False); ws.add_data_validation(dv)
    ws["A1"].value = f"{name}: price breakdown for one {unit}"; ws["A1"].font = H1
    ws["A2"].value = f"Example piece: {example}" if not example.startswith("Example") else example; ws["A2"].font = NOTE
    def cur(row, default="THB"): c = inp(ws, f"E{row}", default); dv.add(c); return c
    conv = lambda amt, row: f"=IF(E{row}=\"AUD\",{amt},{amt}*{REF['fx']})"
    # materials
    band(ws, 4, "1. Materials"); 
    for col, h in zip("BCDEF", ["Qty", "Unit", "Unit cost", "Currency", "Cost (A$)"]): ws[f"{col}4"].value = h; ws[f"{col}4"].font = H2
    for i, m in enumerate(mats):
        r = 5 + i; inp(ws, f"A{r}", m or None); inp(ws, f"B{r}", None, NUM); inp(ws, f"C{r}"); inp(ws, f"D{r}", None, NUM); cur(r); put(ws, f"F{r}", conv(f"B{r}*D{r}", r), fmt=AUD)
    put(ws, "A17", "Materials total", bold=True); put(ws, "F17", "=SUM(F5:F16)", fmt=AUD, bold=True, fill=GREY)
    # labour
    band(ws, 19, "2. Labour")
    for col, h in zip("BCDEF", ["Hours", "", "Rate per hour", "Currency", "Cost (A$)"]): ws[f"{col}19"].value = h; ws[f"{col}19"].font = H2
    for i, t in enumerate(TASKS):
        r = 20 + i; inp(ws, f"A{r}", t); inp(ws, f"B{r}", None, NUM); put(ws, f"C{r}", None)
        c = put(ws, f"D{r}", f"={REF['rate']}", font=GREEN, fmt=NUM); cur(r); put(ws, f"F{r}", conv(f"B{r}*D{r}", r), fmt=AUD)
    put(ws, "A26", "Labour total", bold=True); put(ws, "B26", "=SUM(B20:B25)", fmt=NUM, bold=True, fill=GREY); put(ws, "F26", "=SUM(F20:F25)", fmt=AUD, bold=True, fill=GREY)
    ws["H20"].value = "Rate comes from Settings. Type over it only where a task is paid differently."; ws["H20"].font = NOTE
    # overhead
    band(ws, 28, "3. Overhead, rejects and packaging")
    put(ws, "A29", "Factory overhead per piece"); put(ws, "F29", f"={REF['ohpp']}", font=GREEN, fmt=AUD)
    put(ws, "A30", "Rejects and rework allowance"); put(ws, "D30", f"={REF['rej']}", font=GREEN, fmt=PCT); put(ws, "F30", "=D30*(F17+F26)", fmt=AUD)
    put(ws, "A31", "Packaging (box, bag, card, label)"); inp(ws, "D31", None, NUM); cur(31); put(ws, "F31", conv("D31", 31), fmt=AUD)
    put(ws, "A32", "FACTORY COST PER PIECE", bold=True); put(ws, "F32", "=F17+F26+F29+F30+F31", fmt=AUD, bold=True, fill=RUST)
    # shipping
    def route(top, title, legs):
        band(ws, top, title)
        for col, h in zip("DEF", ["Amount or %", "Currency", "Cost (A$)"]): ws[f"{col}{top}"].value = h; ws[f"{col}{top}"].font = H2
        put(ws, f"A{top+1}", legs[0]); inp(ws, f"D{top+1}", None, NUM); cur(top+1); put(ws, f"F{top+1}", conv(f"D{top+1}", top+1), fmt=AUD)
        put(ws, f"A{top+2}", "Import duty (% of factory cost plus freight)"); inp(ws, f"D{top+2}", None, PCT); put(ws, f"F{top+2}", f"=D{top+2}*(F32+F{top+1})", fmt=AUD)
        put(ws, f"A{top+3}", "Import tax not recovered (% of cost, freight and duty)"); inp(ws, f"D{top+3}", None, PCT); put(ws, f"F{top+3}", f"=D{top+3}*(F32+F{top+1}+F{top+2})", fmt=AUD)
        put(ws, f"A{top+4}", legs[1]); inp(ws, f"D{top+4}", None, NUM); cur(top+4, "AUD"); put(ws, f"F{top+4}", conv(f"D{top+4}", top+4), fmt=AUD)
        put(ws, f"A{top+5}", "Shipping and import total", bold=True); put(ws, f"F{top+5}", f"=SUM(F{top+1}:F{top+4})", fmt=AUD, bold=True, fill=GREY)
    route(34, "4. Shipping, route A: via our address (how it works today)", ["Freight, factory to our address, per piece", "Courier, our address to the customer"])
    route(41, "5. Shipping, route B: factory direct to the customer", ["Freight, factory to the customer, per piece", "Any local handling or delivery charge"])
    ws["H36"].value = "Leave duty and tax at 0% where none is paid, or where the customer pays it."; ws["H36"].font = NOTE
    ws["H35"].value = "For a shared shipment, divide the freight bill by the pieces in it."; ws["H35"].font = NOTE
    # selling
    band(ws, 48, "6. Selling price")
    put(ws, "A49", "Selling price to the customer, GST included (A$)", bold=True); inp(ws, "F49", None, AUD)
    put(ws, "A50", "GST inside that price"); put(ws, "F50", f"=F49-F49/(1+{REF['gst']})", fmt=AUD)
    put(ws, "A51", "Price before GST", bold=True); put(ws, "F51", "=F49-F50", fmt=AUD, bold=True)
    put(ws, "A52", "Shipping charged to the customer on top (A$)"); inp(ws, "F52", None, AUD)
    put(ws, "A53", "Payment fee"); put(ws, "F53", f"=IF(F49>0,(F49+F52)*{REF['fee']}+{REF['fix']},0)", fmt=AUD)
    put(ws, "A54", "Michael's share"); put(ws, "F54", f"=F51*{REF['share']}", fmt=AUD)
    ws["H52"].value = "Type 0 if shipping is included in the price."; ws["H52"].font = NOTE
    band(ws, 56, "7. What Sanchez keeps")
    for col, h in zip("EF", ["Route A", "Route B"]): ws[f"{col}56"].value = h; ws[f"{col}56"].font = H2
    put(ws, "A57", "Landed cost (factory cost plus shipping and import)"); put(ws, "E57", "=F32+F39", fmt=AUD); put(ws, "F57", "=F32+F46", fmt=AUD)
    put(ws, "A58", "PROFIT PER PIECE", bold=True)
    for col in "EF": put(ws, f"{col}58", f"=IF($F$49>0,$F$51+$F$52-$F$53-$F$54-{col}57,0)", fmt=AUD, bold=True, fill=RUST)
    put(ws, "A59", "Margin (profit as a share of the price before GST)", bold=True)
    for col in "EF": put(ws, f"{col}59", f"=IF($F$51>0,{col}58/$F$51,0)", fmt=PCT, bold=True)
    den = lambda m: f"(1-{REF['share']}-{REF['fee']}*(1+{REF['gst']})-{m})"
    put(ws, "A60", "Break-even price, GST included")
    put(ws, "A61", "Price to hit the target margin, GST included")
    for col in "EF":
        num = f"({col}57+{REF['fix']}-$F$52*(1-{REF['fee']}))"
        put(ws, f"{col}60", f"=IF({den('0')}>0,{num}/{den('0')}*(1+{REF['gst']}),0)", fmt=AUD)
        put(ws, f"{col}61", f"=IF({den(REF['target'])}>0,{num}/{den(REF['target'])}*(1+{REF['gst']}),0)", fmt=AUD)
    # what moves the price
    band(ws, 63, "8. What moves the profit (route A)")
    for col, h in zip("EF", ["Profit", "Change"]): ws[f"{col}63"].value = h; ws[f"{col}63"].font = H2
    thb = "(SUMIF(E5:E16,\"THB\",F5:F16)+SUMIF(E20:E25,\"THB\",F20:F25)+IF(E31=\"THB\",F31,0)+IF(E35=\"THB\",F35,0)+IF(E38=\"THB\",F38,0)+F29)"
    sens = [("Materials cost 10% more", "F17*0.1*(1+D30)"), ("Labour hours 10% more", "F26*0.1*(1+D30)"), ("Shipping and import 10% more", "F39*0.1"),
            ("Thai baht 10% stronger against the A$", f"{thb}*0.1"), ("Selling price 10% lower", f"F51*0.1*(1-{REF['share']})-F49*0.1*{REF['fee']}")]
    for i, (lab, d) in enumerate(sens):
        r = 64 + i; put(ws, f"A{r}", lab); put(ws, f"F{r}", f"=IF($F$49>0,-({d}),0)", fmt=AUD); put(ws, f"E{r}", f"=IF($F$49>0,E58+F{r},0)", fmt=AUD)
    ws["H64"].value = "Each line changes one thing and leaves the rest as typed. The duty knock-on is ignored."; ws["H64"].font = NOTE
    ws.freeze_panes = "A4"
    cells[name] = ws

# ---------------- Summary
sm = wb.create_sheet("Summary", 1); widths(sm, {"A": 16, "B": 46, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 12, "I": 14, "J": 12, "K": 15, "L": 15})
sm["A1"].value = "Summary: one example piece per product"; sm["A1"].font = H1
sm["A2"].value = "All figures in A$, carried from the product sheets. Nothing is typed here."; sm["A2"].font = NOTE
heads = ["Product", "Example piece", "Factory cost", "Shipping, route A", "Shipping, route B", "Selling price inc GST", "Profit, route A", "Margin, route A", "Profit, route B", "Margin, route B", "Break-even price, route A", "Target-margin price, route A"]
for i, h in enumerate(heads):
    c = put(sm, f"{'ABCDEFGHIJKL'[i]}4", h, fill=DARK, align="center"); c.font = H2
sm.row_dimensions[4].height = 44
for i, (name, example, unit, _) in enumerate(PRODUCTS):
    r = 5 + i; q = f"'{name}'!"
    put(sm, f"A{r}", name, bold=True); put(sm, f"B{r}", example.replace("Example piece: ", ""))
    for col, ref, fmt in [("C", "F32", AUD), ("D", "F39", AUD), ("E", "F46", AUD), ("F", "F49", AUD), ("G", "E58", AUD), ("H", "E59", PCT), ("I", "F58", AUD), ("J", "F59", PCT), ("K", "E60", AUD), ("L", "E61", AUD)]:
        put(sm, f"{col}{r}", f"={q}{ref}", font=GREEN, fmt=fmt)
sm.freeze_panes = "C5"

if FILL_TEST:  # numbers for checking the formulas only; never shipped
    for k, v in dict(fx=0.045, rate=100, oh=60000, pcs=120, rej=0.05, gst=0.10, fee=0.0175, fix=0.30, share=0.35, target=0.30).items():
        st[REF[k][9:].replace("$", "")].value = v
    ws = cells["Focus mitts"]
    ws["B5"], ws["D5"] = 0.4, 1200; ws["B6"], ws["D6"] = 0.2, 900; ws["B7"], ws["D7"], ws["E7"] = 1, 6, "AUD"
    ws["B20"], ws["B21"] = 1.5, 3
    ws["D31"] = 80; ws["D35"] = 400; ws["D36"] = 0.05; ws["D37"] = 0.10; ws["D38"] = 15
    ws["D42"] = 900; ws["F49"] = 330; ws["F52"] = 20
wb.save(OUT); print("saved", OUT)
